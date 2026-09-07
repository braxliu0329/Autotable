import sys
import os
import json
import logging

# Add parent directory to sys.path to import Autotable modules
sys.path.append(os.path.dirname(os.path.dirname(os.path.abspath(__file__))))

from autotable import AutoTable
from extraction import extract_structure_using_aspose, extract_text_using_aspose
from docx import Document

# Setup logging
logging.basicConfig(level=logging.INFO)
logger = logging.getLogger(__name__)

def create_test_docx(filename):
    doc = Document()
    doc.add_heading('Test Document', 0)
    
    # Create a table with merged cells
    # 2x2 table
    table = doc.add_table(rows=2, cols=2)
    
    # Row 0
    cell_0_0 = table.cell(0, 0)
    cell_0_0.text = "Label:"
    
    cell_0_1 = table.cell(0, 1)
    cell_0_1.text = "_" * 5 # Slot
    
    # Row 1: Merge cells (1,0) and (1,1) -> 1 cell
    cell_1_0 = table.cell(1, 0)
    cell_1_1 = table.cell(1, 1)
    cell_1_0.merge(cell_1_1)
    cell_1_0.text = "Description: " + ("_" * 10) # Slot
    
    doc.save(filename)
    return filename

def test_anchor_injection():
    # Make sure tests directory exists
    if not os.path.exists("tests"):
        os.makedirs("tests")
        
    test_file = os.path.abspath("tests/test_merged_table.docx")
    create_test_docx(test_file)
    print(f"Created test file: {test_file}")
    
    try:
        # 1. Extract Structure using Aspose (Phase 2 feature)
        print("Running Aspose extraction...")
        structure = extract_structure_using_aspose(test_file)
        
        # Check structure format
        # structure is the parsed JSON
        # It should have structure['data']['tables']
        
        if not structure:
            print("Failed: Structure is empty")
            return

        if 'data' not in structure:
             # Maybe structure IS the data if I unpacked it? 
             # No, extraction.py returns json.loads(stdout)
             print(f"Structure keys: {structure.keys()}")
        
        tables = structure.get('data', {}).get('tables', [])
        if not tables:
            print("Failed to extract tables structure from Aspose output")
            return
            
        html_table = tables[0]
        print(f"Extracted HTML Table (First 200 chars):\n{html_table[:200]}...")
        
        # 2. Generate Anchor Plan using python-docx
        print("Generating anchor plan...")
        at = AutoTable("dummy_kb", test_file, None)
        at.load_template()
        
        table = at.doc.tables[0]
        anchor_plan, anchor_map, id_map = at._generate_anchor_plan(table)
        
        print(f"Anchor Plan: {anchor_plan}")
        
        # 3. Inject Anchors into HTML
        print("Injecting anchors...")
        injected_html = at._inject_anchors_to_html(html_table, anchor_plan)
        
        print(f"Injected HTML:\n{injected_html}")
        
        # 4. Verify
        # We expect ID_001 in (0, 1)
        # We expect ID_002 in (1, 0) which is a merged cell
        
        success = True
        if "ID_001" in injected_html and "<strong>" in injected_html:
             print("SUCCESS: ID_001 found in HTML with bold tag")
        else:
             print("FAILURE: ID_001 not found or not bold")
             success = False
             
        if "ID_002" in injected_html:
             print("SUCCESS: ID_002 found in HTML")
        else:
             print("FAILURE: ID_002 not found")
             success = False

        if "colspan" in injected_html:
             print("SUCCESS: HTML contains colspan attributes (Aspose detected merge)")
        else:
             print("WARNING: HTML might not have captured span attributes (or Aspose output format differs)")
             # Note: Aspose might output colspan="2"
             
        # 5. Verify Text Extraction (Legacy/KB path)
        print("\nTesting Text Extraction (KB Path)...")
        text_lines = extract_text_using_aspose(test_file)
        print(f"Extracted {len(text_lines)} lines.")
        if any("Label:" in line for line in text_lines):
            print("SUCCESS: 'Label:' found in extracted text")
        else:
            print("FAILURE: 'Label:' not found in extracted text")
            print(f"Content: {text_lines}")

        if success:
            print("\n=== PHASE 2 TEST PASSED ===")
        else:
            print("\n=== PHASE 2 TEST FAILED ===")

    except Exception as e:
        print(f"Test failed with error: {e}")
        import traceback
        traceback.print_exc()

if __name__ == "__main__":
    test_anchor_injection()
